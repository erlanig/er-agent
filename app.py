import csv
import hashlib
import hmac
import io
import json
import os
import re
import time
import zipfile
import xml.etree.ElementTree as ET

from dotenv import load_dotenv

load_dotenv()  # baca .env saat lokal; di Vercel dipakai Environment Variables

from flask import Flask, Response, abort, jsonify, render_template, request, send_from_directory, stream_with_context  # noqa: E402

import approvals  # noqa: E402
import db  # noqa: E402
from providers import MODELS, diagnose, run_chat  # noqa: E402
from tools import google_status, integrations, tool_specs  # noqa: E402
from tools_extra import run_due_jobs  # noqa: E402
import legal  # noqa: E402

app = Flask(__name__)
legal.register(app)  # /about, /privacy, /google<code>.html

MAX_MESSAGES = 40
MAX_CHARS = 80_000
MAX_INSTRUCTIONS = 6_000
MAX_CHAT_BYTES = 1_500_000  # batas Vercel untuk body request adalah 4,5 MB
SESSION_COOKIE = "bc_session"
SESSION_DAYS = 30
_ID_RE = re.compile(r"^[\w-]{1,40}$")
_CHAT_KEYS = ("id", "title", "model", "messages", "pinned", "projectId", "updated")


# --------------------------------------------------------------------------
# Auth: password sekali, lalu cookie sesi bertanda tangan (30 hari)
# --------------------------------------------------------------------------
def _secret() -> bytes:
    return (os.getenv("SESSION_SECRET") or os.getenv("APP_PASSWORD") or "").encode()


def _sign(exp: int) -> str:
    return hmac.new(_secret(), str(exp).encode(), hashlib.sha256).hexdigest()


def _make_token() -> str:
    exp = int(time.time()) + SESSION_DAYS * 86400
    return f"{exp}.{_sign(exp)}"


def _token_ok(token: str) -> bool:
    try:
        exp_s, sig = token.split(".", 1)
        exp = int(exp_s)
    except (ValueError, AttributeError):
        return False
    return exp > time.time() and hmac.compare_digest(sig, _sign(exp))


def _password_ok(pw: str) -> bool:
    expected = os.getenv("APP_PASSWORD", "")
    return bool(expected) and hmac.compare_digest((pw or "").encode(), expected.encode())


def _auth_state() -> str:
    """'ok' | 'denied' | 'unconfigured'"""
    if not os.getenv("APP_PASSWORD"):
        return "unconfigured"
    if _token_ok(request.cookies.get(SESSION_COOKIE, "")):
        return "ok"
    if _password_ok(request.headers.get("X-App-Password", "")):
        return "ok"
    return "denied"


def _auth_error(state: str):
    if state == "unconfigured":
        return jsonify(error="APP_PASSWORD belum diset di environment."), 503
    return jsonify(error="Sesi berakhir atau password salah."), 401


def _is_https() -> bool:
    return request.is_secure or request.headers.get("X-Forwarded-Proto", "").startswith("https")


# --------------------------------------------------------------------------
# Ekstraktor Dokumen (PDF, DOCX, XLSX, CSV, Teks)
# --------------------------------------------------------------------------
def _extract_docx(raw: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        xml_content = z.read("word/document.xml")
    tree = ET.fromstring(xml_content)
    paragraphs = []
    for p in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
        texts = [t.text for t in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t") if t.text]
        if texts:
            paragraphs.append("".join(texts))
    return "\n".join(paragraphs)


def _extract_pdf(raw: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(raw))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages[:50])
    except ImportError:
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(io.BytesIO(raw))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages[:50])
        except ImportError:
            raise RuntimeError("Library 'pypdf' belum terinstal di server (jalankan: pip install pypdf).")


def _extract_xlsx(raw: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        out = []
        for sheet in wb.sheetnames[:10]:
            ws = wb[sheet]
            rows = []
            for idx, row in enumerate(ws.iter_rows(values_only=True)):
                if idx > 300:
                    rows.append("... (baris berikutnya dipotong)")
                    break
                if any(cell is not None for cell in row):
                    rows.append("\t".join("" if c is None else str(c) for c in row))
            if rows:
                out.append(f"### Sheet: {sheet}\n" + "\n".join(rows))
        return "\n\n".join(out)
    except ImportError:
        raise RuntimeError("Library 'openpyxl' belum terinstal di server (jalankan: pip install openpyxl).")


# --------------------------------------------------------------------------
# Route
# --------------------------------------------------------------------------
@app.get("/")
def index():
    return render_template("index.html")


def _asset(name: str):
    """Sajikan icon dari static/ atau root proyek (publik, tanpa auth agar tampil di halaman login)."""
    for folder in (os.path.join(app.root_path, "static"), app.root_path):
        if os.path.isfile(os.path.join(folder, name)):
            return send_from_directory(folder, name, max_age=86400)
    abort(404)


@app.get("/bird.png")
def bird_png():
    return _asset("bird.png")


@app.get("/bird.ico")
@app.get("/favicon.ico")
def bird_ico():
    return _asset("bird.ico")


@app.get("/api/health")
def health():
    return jsonify(ok=True)


@app.get("/api/config")
def config():
    return jsonify(
        models=[{"key": k, "label": v["label"], "provider": v["provider"]} for k, v in MODELS.items()],
        tools=[
            {"name": t["name"], "description": t["description"], "dangerous": t["dangerous"],
             "group": t["group"], "available": t["available"]}
            for t in tool_specs(allow_actions=True)
        ],
        integrations=integrations(),
        features={"approval": True, "sync": db.enabled()},  # sync = riwayat disimpan di database server
    )


@app.get("/api/google/status")
def google_status_route():
    state = _auth_state()
    if state != "ok":
        return _auth_error(state)
    try:
        return jsonify(google_status())
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 500


@app.get("/api/session")
def session():
    state = _auth_state()
    return jsonify(ok=state == "ok", configured=state != "unconfigured")


@app.post("/api/login")
def login():
    if not os.getenv("APP_PASSWORD"):
        return _auth_error("unconfigured")
    body = request.get_json(silent=True) or {}
    pw = body.get("password") or request.headers.get("X-App-Password", "")
    if not _password_ok(pw):
        time.sleep(0.4)
        return jsonify(error="Password salah."), 401
    resp = jsonify(ok=True)
    resp.set_cookie(
        SESSION_COOKIE, _make_token(), max_age=SESSION_DAYS * 86400,
        httponly=True, secure=_is_https(), samesite="Lax", path="/",
    )
    return resp


@app.post("/api/logout")
def logout():
    resp = jsonify(ok=True)
    resp.delete_cookie(SESSION_COOKIE, path="/")
    return resp


@app.post("/api/extract")
def extract_file():
    state = _auth_state()
    if state != "ok":
        return _auth_error(state)
    f = request.files.get("file")
    if not f:
        return jsonify(error="Tidak ada file yang dikirim."), 400

    filename = (f.filename or "").lower()
    raw = f.read(15 * 1024 * 1024)
    try:
        if filename.endswith(".docx"):
            text = _extract_docx(raw)
        elif filename.endswith(".pdf"):
            text = _extract_pdf(raw)
        elif filename.endswith(".xlsx"):
            text = _extract_xlsx(raw)
        elif filename.endswith(".csv"):
            decoded = raw.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(decoded))
            lines = ["\t".join(r) for _, r in zip(range(500), reader)]
            text = "\n".join(lines)
        else:
            text = raw.decode("utf-8", errors="replace")
        return jsonify(ok=True, text=text[:MAX_CHARS])
    except Exception as e:
        return jsonify(error=str(e)), 422


@app.post("/api/diagnose")
def diagnose_route():
    state = _auth_state()
    if state != "ok":
        return _auth_error(state)
    try:
        return jsonify(results=diagnose())
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 500


@app.post("/api/action")
def action_route():
    """Jalankan aksi yang sudah disetujui pengguna lewat kartu persetujuan."""
    state = _auth_state()
    if state != "ok":
        return _auth_error(state)
    body = request.get_json(silent=True) or {}
    try:
        out = approvals.run_approved(
            str(body.get("name") or ""), body.get("args") or {},
            int(body.get("exp") or 0), str(body.get("token") or ""),
        )
    except PermissionError as e:
        return jsonify(error=str(e)), 403
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 500
    return jsonify(out)


# --------------------------------------------------------------------------
# Cron: kirim email / Telegram terjadwal yang sudah jatuh tempo
# --------------------------------------------------------------------------
@app.route("/api/cron/run-jobs", methods=["GET", "POST"])
def cron_run_jobs():
    """Dipanggil berkala (cron-job.org / GitHub Actions / Vercel Cron) dengan
    header `Authorization: Bearer <CRON_SECRET>` (atau ?key=<CRON_SECRET>)."""
    secret = os.getenv("CRON_SECRET", "")
    auth = request.headers.get("Authorization", "")
    given = auth[7:] if auth.startswith("Bearer ") else request.args.get("key", "")
    if not secret or not hmac.compare_digest(given.encode(), secret.encode()):
        return jsonify(error="Unauthorized"), 401
    try:
        return jsonify(run_due_jobs())
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 500


# --------------------------------------------------------------------------
# Sinkronisasi riwayat (Postgres)
# --------------------------------------------------------------------------
def _sync_guard():
    """None bila boleh lanjut; selain itu response error."""
    state = _auth_state()
    if state != "ok":
        return _auth_error(state)
    if not db.enabled():
        return jsonify(error="DATABASE_URL belum diset.", configured=False), 503
    return None


@app.get("/api/sync")
def sync_pull():
    err = _sync_guard()
    if err:
        return err
    try:
        chats, projects = db.load_all()
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 500
    return jsonify(chats=chats, projects=projects)


@app.put("/api/chats/<chat_id>")
def sync_put_chat(chat_id: str):
    err = _sync_guard()
    if err:
        return err
    if not _ID_RE.match(chat_id):
        return jsonify(error="ID tidak valid."), 400
    raw = request.get_data()
    if len(raw) > MAX_CHAT_BYTES:
        return jsonify(error="Obrolan terlalu besar untuk disimpan di server."), 413
    try:
        data = json.loads(raw)
    except ValueError:
        return jsonify(error="JSON tidak valid."), 400
    if not isinstance(data, dict) or data.get("id") != chat_id or not isinstance(data.get("messages"), list):
        return jsonify(error="Data obrolan tidak valid."), 400
    data = {k: data[k] for k in _CHAT_KEYS if k in data}
    try:
        db.upsert_chat(chat_id, data)
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 500
    return jsonify(ok=True)


@app.delete("/api/chats/<chat_id>")
def sync_delete_chat(chat_id: str):
    err = _sync_guard()
    if err:
        return err
    if not _ID_RE.match(chat_id):
        return jsonify(error="ID tidak valid."), 400
    try:
        db.delete_chat(chat_id)
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 500
    return jsonify(ok=True)


@app.put("/api/projects")
def sync_put_projects():
    err = _sync_guard()
    if err:
        return err
    body = request.get_json(silent=True)
    if not isinstance(body, list) or len(body) > 200:
        return jsonify(error="Daftar proyek tidak valid."), 400
    projects = [
        {"id": str(p.get("id", ""))[:40], "name": str(p.get("name", ""))[:120],
         "instructions": str(p.get("instructions", ""))[:6000], "open": bool(p.get("open"))}
        for p in body if isinstance(p, dict) and _ID_RE.match(str(p.get("id", "")))
    ]
    try:
        db.set_projects(projects)
    except Exception as e:
        return jsonify(error=f"{type(e).__name__}: {e}"), 500
    return jsonify(ok=True)


@app.post("/api/chat")
def chat():
    state = _auth_state()
    if state != "ok":
        return _auth_error(state)

    body = request.get_json(silent=True) or {}
    model_key = body.get("model")
    if model_key not in MODELS:
        return jsonify(error="Model tidak dikenal."), 400

    history = []
    for m in (body.get("messages") or [])[-MAX_MESSAGES:]:
        role, content = m.get("role"), m.get("content")
        images = m.get("images") or []
        if role in ("user", "assistant") and isinstance(content, str) and (content.strip() or images):
            clean_imgs = []
            for img in images[:4]:
                if isinstance(img, dict) and str(img.get("data_url", "")).startswith("data:image/"):
                    clean_imgs.append({
                        "media_type": str(img.get("media_type") or "image/png"),
                        "data_url": str(img["data_url"]),
                    })
            history.append({"role": role, "content": content[:MAX_CHARS], "images": clean_imgs})

    while history and history[0]["role"] != "user":
        history.pop(0)
    if not history or history[-1]["role"] != "user":
        return jsonify(error="Pesan terakhir harus dari pengguna."), 400

    allow_actions = bool(body.get("allow_actions"))
    approval = bool(body.get("approval_mode"))  # UI baru mengirim true -> aksi menunggu tombol Setujui
    enabled = body.get("enabled_tools")
    enabled = {str(x) for x in enabled} if isinstance(enabled, list) else None
    extra = str(body.get("instructions") or "")[:MAX_INSTRUCTIONS]

    def generate():
        try:
            for event in run_chat(model_key, history, allow_actions, enabled, extra, approval):
                yield json.dumps(event, ensure_ascii=False) + "\n"
        except Exception as e:
            yield json.dumps({"type": "error", "message": f"{type(e).__name__}: {e}"}, ensure_ascii=False) + "\n"
        yield json.dumps({"type": "done"}) + "\n"

    return Response(
        stream_with_context(generate()),
        mimetype="application/x-ndjson",
        headers={"Cache-Control": "no-cache, no-transform", "X-Accel-Buffering": "no"},
    )


if __name__ == "__main__":
    app.run(debug=True, port=int(os.getenv("PORT", "5000")))